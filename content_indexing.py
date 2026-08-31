from __future__ import annotations

import html
import ipaddress
import json
import os
import re
import socket
import subprocess
import tempfile
import wave
from dataclasses import dataclass
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from PIL import Image, ImageOps

from tgbrain import (
    ENRICHMENT_VERSION,
    METADATA_VERSION,
    Settings,
    VIDEO_VISION_PROMPT_VERSION,
    VISION_PROMPT_VERSION,
    build_indexed_text,
    describe,
    embed,
    message_metadata,
)


ROOT = Path(__file__).resolve().parent
VISION_SOURCE = ROOT / "tools" / "vision_extract.swift"
VISION_BINARY = ROOT / "data" / "bin" / "vision-extract"
MAX_EXTRACTED_CHARS = 500_000
MAX_LINK_BYTES = 2_000_000
VIDEO_FRAME_COUNT = 6
VIDEO_FRAME_SIZE = (480, 300)
VIDEO_VISION_PROMPT = (
    "These are representative frames sampled in time order from one saved "
    "video. Index the video for a private visual search library. Return "
    "concise plain text using exactly these labels: Summary, People, Visible "
    "text, Objects, Setting, Sequence, Meme context, Search terms. Describe "
    "what happens across the frames, not each frame separately. Read useful "
    "on-screen text, but keep Visible text to the 15 most meaningful labels. "
    "Name recognizable public figures only when confident. "
    "For a meme or screen recording, identify the subject, app or site, action "
    "and joke. Search terms must include concrete names, entities, actions, "
    "visual traits, topics and likely user query phrases. Use Unknown rather "
    "than inventing identity. Stay below 220 words and do not add a preamble."
)


class ContentUnavailable(RuntimeError):
    pass


class MetadataParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title_parts: list[str] = []
        self.in_title = False
        self.metadata: dict[str, str] = {}

    def handle_starttag(
        self,
        tag: str,
        attrs: list[tuple[str, str | None]],
    ) -> None:
        attributes = {
            key.lower(): (value or "").strip()
            for key, value in attrs
        }
        if tag.lower() == "title":
            self.in_title = True
        if tag.lower() != "meta":
            return
        key = (
            attributes.get("property")
            or attributes.get("name")
            or ""
        ).lower()
        value = attributes.get("content") or ""
        if key and value:
            self.metadata[key] = value

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "title":
            self.in_title = False

    def handle_data(self, data: str) -> None:
        if self.in_title and data.strip():
            self.title_parts.append(data.strip())

    def result(self) -> dict[str, str]:
        title = (
            self.metadata.get("og:title")
            or self.metadata.get("twitter:title")
            or " ".join(self.title_parts)
        )
        description = (
            self.metadata.get("og:description")
            or self.metadata.get("twitter:description")
            or self.metadata.get("description")
            or ""
        )
        site_name = self.metadata.get("og:site_name") or ""
        return {
            "title": html.unescape(title).strip()[:500],
            "description": html.unescape(description).strip()[:1200],
            "site_name": html.unescape(site_name).strip()[:200],
        }


@dataclass
class EnrichmentResult:
    row_id: int
    vision_text: str
    vision_model: str
    vision_prompt_version: int
    vision_attempted_at: str | None
    extracted_text: str
    link_metadata_json: str
    content_status: str
    content_error: str
    indexed_text: str
    embedding_json: str | None
    embedding_model: str
    embedding_attempted_at: str | None
    category: str
    is_sensitive: int
    sensitive_reason: str | None


def _clean_text(value: str) -> str:
    value = value.replace("\x00", " ").strip()
    return value[:MAX_EXTRACTED_CHARS]


def _valid_vision_description(value: str) -> bool:
    words = re.findall(r"[^\W_][\w'-]+", value, re.UNICODE)
    lowered = value.lower()
    return (
        len(value.strip()) >= 45
        and len(words) >= 8
        and not lowered.startswith(("got it", "let's break", "we need"))
    )


def ensure_vision_binary() -> Path:
    if (
        VISION_BINARY.exists()
        and VISION_BINARY.stat().st_mtime >= VISION_SOURCE.stat().st_mtime
    ):
        return VISION_BINARY
    VISION_BINARY.parent.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [
            "/usr/bin/swiftc",
            "-O",
            str(VISION_SOURCE),
            "-o",
            str(VISION_BINARY),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    if completed.returncode != 0:
        raise ContentUnavailable(
            completed.stderr.strip() or "macOS Vision compiler failed"
        )
    return VISION_BINARY


def extract_with_vision(path: Path) -> str:
    binary = ensure_vision_binary()
    completed = subprocess.run(
        [str(binary), str(path)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=240,
        check=False,
    )
    if completed.returncode != 0:
        raise ContentUnavailable(
            completed.stderr.strip() or "macOS Vision extraction failed"
        )
    return _clean_text(completed.stdout)


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.append("\t".join(cell.text for cell in row.cells))
    return _clean_text("\n".join(value for value in values if value.strip()))


def _extract_presentation(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(path)
    values = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_values = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                slide_values.append(text.strip())
        if slide_values:
            values.append(
                f"Slide {slide_number}\n" + "\n".join(slide_values)
            )
    return _clean_text("\n\n".join(values))


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    values = []
    for sheet in workbook.worksheets:
        values.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            rendered = "\t".join(
                str(value) for value in row if value is not None
            )
            if rendered:
                values.append(rendered)
    workbook.close()
    return _clean_text("\n".join(values))


def _extract_xls(path: Path) -> str:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    values = []
    for sheet in workbook.sheets():
        values.append(f"Sheet: {sheet.name}")
        for row_index in range(sheet.nrows):
            rendered = "\t".join(
                str(sheet.cell_value(row_index, column_index))
                for column_index in range(sheet.ncols)
                if sheet.cell_value(row_index, column_index) not in {"", None}
            )
            if rendered:
                values.append(rendered)
    workbook.release_resources()
    return _clean_text("\n".join(values))


def _extract_pdf_fallback(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return _clean_text(
        "\n\n".join((page.extract_text() or "") for page in reader.pages)
    )


def extract_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            return extract_with_vision(path)
        except ContentUnavailable:
            return _extract_pdf_fallback(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pptx":
        return _extract_presentation(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(path)
    if suffix == ".xls":
        return _extract_xls(path)
    if suffix in {
        ".txt",
        ".md",
        ".csv",
        ".tsv",
        ".json",
        ".xml",
        ".html",
        ".htm",
        ".log",
        ".py",
        ".js",
        ".ts",
        ".css",
    }:
        return _clean_text(
            path.read_text(encoding="utf-8", errors="replace")
        )
    raise ContentUnavailable(
        f"No local text extractor for {suffix or 'this document type'}"
    )


def _whisper_paths(config: Settings) -> tuple[Path, Path]:
    command = (
        Path(config.whisper_command).expanduser()
        if config.whisper_command
        else ROOT
        / "data"
        / "tools"
        / "whisper.cpp"
        / "build"
        / "bin"
        / "whisper-cli"
    )
    model = config.whisper_model_path or (
        ROOT / "data" / "models" / "ggml-small.bin"
    )
    if not command.exists():
        raise ContentUnavailable(
            "Local voice model is not installed; run "
            "./setup_content_tools.sh"
        )
    if not model.exists():
        raise ContentUnavailable(
            "Local voice model is missing; run ./setup_content_tools.sh"
        )
    return command.resolve(), model.resolve()


def transcribe_audio(config: Settings, path: Path) -> str:
    command, model = _whisper_paths(config)
    with tempfile.TemporaryDirectory() as directory:
        import av

        wav_path = Path(directory) / "input.wav"
        container = av.open(str(path))
        stream = next(
            (stream for stream in container.streams if stream.type == "audio"),
            None,
        )
        if stream is None:
            container.close()
            raise ContentUnavailable("The media file has no audio stream")
        resampler = av.AudioResampler(
            format="s16",
            layout="mono",
            rate=16000,
        )
        with wave.open(str(wav_path), "wb") as output_wave:
            output_wave.setnchannels(1)
            output_wave.setsampwidth(2)
            output_wave.setframerate(16000)
            for frame in container.decode(stream):
                converted = resampler.resample(frame)
                for output_frame in (
                    converted if isinstance(converted, list) else [converted]
                ):
                    if output_frame is not None:
                        output_wave.writeframes(
                            output_frame.to_ndarray().tobytes()
                        )
            tail = resampler.resample(None)
            for output_frame in tail if isinstance(tail, list) else [tail]:
                if output_frame is not None:
                    output_wave.writeframes(
                        output_frame.to_ndarray().tobytes()
                    )
        container.close()

        output = Path(directory) / "transcript"
        completed = subprocess.run(
            [
                str(command),
                "-m",
                str(model),
                "-f",
                str(wav_path),
                "-otxt",
                "-of",
                str(output),
                "-np",
                "-nt",
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            timeout=1200,
            check=False,
        )
        transcript = output.with_suffix(".txt")
        if completed.returncode != 0 or not transcript.exists():
            raise ContentUnavailable(
                completed.stderr.strip()
                or completed.stdout.strip()
                or "Local voice transcription failed"
            )
        return _clean_text(
            transcript.read_text(encoding="utf-8", errors="replace")
        )


def video_has_audio(path: Path) -> bool:
    import av

    with av.open(str(path)) as container:
        return any(stream.type == "audio" for stream in container.streams)


def sample_video_frames(path: Path, count: int = VIDEO_FRAME_COUNT) -> list[Image.Image]:
    import av

    frames: list[Image.Image] = []
    fingerprints: set[bytes] = set()
    with av.open(str(path)) as container:
        stream = next(
            (candidate for candidate in container.streams if candidate.type == "video"),
            None,
        )
        if stream is None:
            raise ContentUnavailable("The media file has no video stream")
        duration = (
            float(container.duration / av.time_base)
            if container.duration
            else 0.0
        )
        if duration > 0:
            if count == 1:
                targets = [duration * 0.5]
            else:
                targets = [
                    duration * (0.06 + (0.88 * index / (count - 1)))
                    for index in range(count)
                ]
        else:
            targets = [float(index) for index in range(count)]

        for target in targets:
            try:
                container.seek(
                    max(0, int(target * av.time_base)),
                    any_frame=False,
                    backward=True,
                )
                frame = next(container.decode(stream), None)
            except Exception:
                frame = None
            if frame is None:
                continue
            image = frame.to_image().convert("RGB")
            fingerprint = ImageOps.fit(
                image,
                (24, 24),
                method=Image.Resampling.BILINEAR,
            ).tobytes()
            if fingerprint in fingerprints:
                continue
            fingerprints.add(fingerprint)
            image.thumbnail(VIDEO_FRAME_SIZE, Image.Resampling.LANCZOS)
            frames.append(image)

    if not frames:
        raise ContentUnavailable("No readable frames were found in the video")
    return frames


def build_video_contact_sheet(path: Path, output_path: Path) -> Path:
    frames = sample_video_frames(path)
    columns = 2
    rows = (len(frames) + columns - 1) // columns
    sheet = Image.new(
        "RGB",
        (VIDEO_FRAME_SIZE[0] * columns, VIDEO_FRAME_SIZE[1] * rows),
        "black",
    )
    for index, frame in enumerate(frames):
        tile = ImageOps.contain(
            frame,
            VIDEO_FRAME_SIZE,
            method=Image.Resampling.LANCZOS,
        )
        left = (index % columns) * VIDEO_FRAME_SIZE[0]
        top = (index // columns) * VIDEO_FRAME_SIZE[1]
        x = left + (VIDEO_FRAME_SIZE[0] - tile.width) // 2
        y = top + (VIDEO_FRAME_SIZE[1] - tile.height) // 2
        sheet.paste(tile, (x, y))
    sheet.save(output_path, format="JPEG", quality=88, optimize=True)
    return output_path


def describe_video(config: Settings, path: Path) -> str:
    with tempfile.TemporaryDirectory() as directory:
        contact_sheet = build_video_contact_sheet(
            path,
            Path(directory) / "video-frames.jpg",
        )
        frame_text = ""
        try:
            frame_text = extract_with_vision(contact_sheet)
        except Exception:
            pass
        description = describe(
            config,
            str(contact_sheet),
            prompt=VIDEO_VISION_PROMPT,
            max_edge=1280,
            max_tokens=400,
            timeout_seconds=180,
        )
        if frame_text and frame_text.lower() not in description.lower():
            description += "\nFrame OCR: " + clean_preview_text(frame_text, 1200)
        return _clean_text(description)


def clean_preview_text(value: str, limit: int) -> str:
    compact = re.sub(r"\s+", " ", value or "").strip()
    return compact if len(compact) <= limit else compact[: limit - 1].rstrip() + "…"


def _safe_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ContentUnavailable("Only public HTTP links are supported")
    if parsed.username or parsed.password:
        raise ContentUnavailable("Credential-bearing links are not fetched")
    try:
        addresses = socket.getaddrinfo(
            parsed.hostname,
            parsed.port or (443 if parsed.scheme == "https" else 80),
            type=socket.SOCK_STREAM,
        )
    except socket.gaierror as error:
        raise ContentUnavailable("Link host could not be resolved") from error
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if not ip.is_global:
            raise ContentUnavailable("Private or local links are not fetched")


def fetch_link_metadata(url: str) -> dict:
    current = url
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Apple Silicon Mac OS X) "
            "AppleWebKit/537.36 TelegramBrain/1.0"
        ),
        "Accept": "text/html,application/xhtml+xml",
    }
    response = None
    for _ in range(5):
        _safe_public_url(current)
        response = requests.get(
            current,
            headers=headers,
            timeout=(4, 10),
            allow_redirects=False,
            stream=True,
        )
        if response.is_redirect or response.is_permanent_redirect:
            target = response.headers.get("location")
            response.close()
            if not target:
                raise ContentUnavailable("Link redirect had no destination")
            current = urljoin(current, target)
            continue
        break
    if response is None:
        raise ContentUnavailable("Link could not be loaded")
    try:
        response.raise_for_status()
        content_type = response.headers.get("content-type", "").lower()
        if "html" not in content_type:
            return {
                "url": url,
                "resolved_url": current,
                "title": Path(urlparse(current).path).name,
                "description": content_type.split(";")[0],
                "site_name": urlparse(current).hostname or "",
                "status": "ready",
            }
        chunks = []
        size = 0
        for chunk in response.iter_content(32_768):
            size += len(chunk)
            if size > MAX_LINK_BYTES:
                break
            chunks.append(chunk)
        encoding = response.encoding or "utf-8"
        source = b"".join(chunks).decode(encoding, errors="replace")
    finally:
        response.close()
    parser = MetadataParser()
    parser.feed(source)
    metadata = parser.result()
    return {
        "url": url,
        "resolved_url": current,
        **metadata,
        "status": "ready",
    }


def _extract_media(config: Settings, row: dict) -> str:
    path = Path(row.get("media_path") or "")
    if not path.exists():
        raise ContentUnavailable("The downloaded local file is missing")
    kind = row.get("media_type")
    if kind == "image":
        if path.suffix.lower() == ".gif":
            raise ContentUnavailable("Animated image OCR is not supported")
        return extract_with_vision(path)
    if kind in {"pdf", "document"}:
        return extract_document(path)
    if kind == "audio":
        return transcribe_audio(config, path)
    if kind == "video":
        return transcribe_audio(config, path) if video_has_audio(path) else ""
    return ""


def enrich_record(config: Settings, row: dict) -> EnrichmentResult:
    errors = []
    successes = 0
    expected = 0
    extracted_text = row.get("extracted_text") or ""
    vision_text = row.get("vision_text") or ""
    vision_model = row.get("vision_model") or ""
    vision_prompt_version = int(row.get("vision_prompt_version") or 0)
    vision_attempted_at = row.get("vision_attempted_at")
    media_path = row.get("media_path")
    if media_path:
        expected += 1
        if int(row.get("enrichment_version") or 0) >= ENRICHMENT_VERSION:
            successes += 1
        else:
            try:
                extracted_text = _extract_media(config, row)
                successes += 1
            except Exception as error:
                errors.append(f"Media: {error}")
        media_type = row.get("media_type")
        if media_type in {"image", "video"} and config.enable_vision:
            expected += 1
            vision_attempted_at = datetime.now(timezone.utc).isoformat()
            try:
                candidate_vision_text = (
                    describe_video(config, Path(media_path))
                    if media_type == "video"
                    else describe(config, media_path)
                )
                if not _valid_vision_description(candidate_vision_text):
                    raise ContentUnavailable(
                        "The vision model returned no useful description"
                    )
                vision_text = candidate_vision_text
                vision_model = config.vision_model
                vision_prompt_version = (
                    VIDEO_VISION_PROMPT_VERSION
                    if media_type == "video"
                    else VISION_PROMPT_VERSION
                )
                successes += 1
            except Exception as error:
                label = "Video understanding" if media_type == "video" else "Image understanding"
                errors.append(f"{label}: {error}")

    try:
        urls = json.loads(row.get("urls_json") or "[]")
    except (TypeError, json.JSONDecodeError):
        urls = []
    link_metadata = []
    for url in urls if isinstance(urls, list) else []:
        if not isinstance(url, str):
            continue
        expected += 1
        if row.get("is_sensitive"):
            link_metadata.append(
                {
                    "url": url,
                    "status": "protected",
                    "error": "Sensitive links are not fetched",
                }
            )
            errors.append("Link metadata protected")
            continue
        try:
            link_metadata.append(fetch_link_metadata(url))
            successes += 1
        except Exception as error:
            link_metadata.append(
                {
                    "url": url,
                    "status": "unavailable",
                    "error": str(error)[:240],
                }
            )
            errors.append(f"Link: {error}")

    if not expected or successes == expected:
        status = "ready"
    elif successes:
        status = "partial"
    else:
        status = "failed"

    updated = {
        **row,
        "vision_text": vision_text,
        "extracted_text": extracted_text,
        "link_metadata_json": json.dumps(link_metadata),
    }
    updated["indexed_text"] = build_indexed_text(updated)
    metadata = message_metadata(updated)
    embedding_json = row.get("embedding_json")
    embedding_model = row.get("embedding_model") or ""
    embedding_attempted_at = row.get("embedding_attempted_at")
    if config.enable_embeddings and updated["indexed_text"]:
        embedding_attempted_at = datetime.now(timezone.utc).isoformat()
        try:
            embedding_json = json.dumps(
                embed(config, updated["indexed_text"])
            )
            embedding_model = config.embed_model
        except Exception as error:
            errors.append(f"Embedding: {error}")
            if status == "ready":
                status = "partial"

    return EnrichmentResult(
        row_id=int(row["id"]),
        vision_text=vision_text,
        vision_model=vision_model,
        vision_prompt_version=vision_prompt_version,
        vision_attempted_at=vision_attempted_at,
        extracted_text=extracted_text,
        link_metadata_json=json.dumps(link_metadata),
        content_status=status,
        content_error=" · ".join(errors)[:2000],
        indexed_text=updated["indexed_text"],
        embedding_json=embedding_json,
        embedding_model=embedding_model,
        embedding_attempted_at=embedding_attempted_at,
        category=metadata["category"],
        is_sensitive=metadata["is_sensitive"],
        sensitive_reason=metadata["sensitive_reason"],
    )


def apply_enrichment(connection, result: EnrichmentResult) -> None:
    connection.execute(
        """
        UPDATE messages
        SET vision_text = ?, vision_model = ?, vision_prompt_version = ?,
            vision_attempted_at = ?, extracted_text = ?, link_metadata_json = ?,
            content_status = ?, content_error = ?,
            enrichment_version = ?, enriched_at = ?,
            indexed_text = ?, embedding_json = ?, embedding_model = ?,
            embedding_attempted_at = ?,
            category = ?, is_sensitive = ?, sensitive_reason = ?,
            metadata_version = ?
        WHERE id = ?
        """,
        (
            result.vision_text,
            result.vision_model,
            result.vision_prompt_version,
            result.vision_attempted_at,
            result.extracted_text,
            result.link_metadata_json,
            result.content_status,
            result.content_error,
            ENRICHMENT_VERSION,
            datetime.now(timezone.utc).isoformat(),
            result.indexed_text,
            result.embedding_json,
            result.embedding_model,
            result.embedding_attempted_at,
            result.category,
            result.is_sensitive,
            result.sensitive_reason,
            METADATA_VERSION,
            result.row_id,
        ),
    )
    connection.commit()


def pending_count(
    connection,
    *,
    require_vision: bool = False,
    vision_model: str = "",
    require_embeddings: bool = False,
    embed_model: str = "",
) -> int:
    missing_vision = (
        """
        OR (
            media_path IS NOT NULL
            AND (
                (
                    media_type = 'image'
                    AND (
                        TRIM(vision_text) = ''
                        OR COALESCE(vision_model, '') != ?
                        OR vision_prompt_version < ?
                    )
                )
                OR (
                    media_type = 'video'
                    AND (
                        TRIM(vision_text) = ''
                        OR COALESCE(vision_model, '') != ?
                        OR vision_prompt_version < ?
                    )
                )
            )
        )
        """
        if require_vision
        else ""
    )
    missing_embedding = (
        """
        OR (
            TRIM(indexed_text) != ''
            AND (
                embedding_json IS NULL
                OR COALESCE(embedding_model, '') != ?
            )
        )
        """
        if require_embeddings
        else ""
    )
    params: list = [ENRICHMENT_VERSION]
    if require_vision:
        params.extend(
            [
                vision_model,
                VISION_PROMPT_VERSION,
                vision_model,
                VIDEO_VISION_PROMPT_VERSION,
            ]
        )
    if require_embeddings:
        params.append(embed_model)
    return int(
        connection.execute(
            f"""
            SELECT COUNT(*)
            FROM messages
            WHERE duplicate_of_id IS NULL
              AND (
                  enrichment_version < ?
                  OR content_status IN ('pending', 'processing')
                  {missing_vision}
                  {missing_embedding}
              )
            """,
            params,
        ).fetchone()[0]
    )


def next_pending(
    connection,
    limit: int = 12,
    *,
    require_vision: bool = False,
    vision_model: str = "",
    require_embeddings: bool = False,
    embed_model: str = "",
) -> list[dict]:
    missing_vision = (
        """
        OR (
            media_path IS NOT NULL
            AND (
                (
                    media_type = 'image'
                    AND (
                        TRIM(vision_text) = ''
                        OR COALESCE(vision_model, '') != ?
                        OR vision_prompt_version < ?
                    )
                )
                OR (
                    media_type = 'video'
                    AND (
                        TRIM(vision_text) = ''
                        OR COALESCE(vision_model, '') != ?
                        OR vision_prompt_version < ?
                    )
                )
            )
            AND (
                vision_attempted_at IS NULL
                OR datetime(vision_attempted_at) <= datetime('now', '-6 hours')
            )
        )
        """
        if require_vision
        else ""
    )
    missing_embedding = (
        """
        OR (
            TRIM(indexed_text) != ''
            AND (
                embedding_json IS NULL
                OR COALESCE(embedding_model, '') != ?
            )
            AND (
                embedding_attempted_at IS NULL
                OR datetime(embedding_attempted_at)
                    <= datetime('now', '-6 hours')
            )
        )
        """
        if require_embeddings
        else ""
    )
    params: list = [ENRICHMENT_VERSION]
    if require_vision:
        params.extend(
            [
                vision_model,
                VISION_PROMPT_VERSION,
                vision_model,
                VIDEO_VISION_PROMPT_VERSION,
            ]
        )
    if require_embeddings:
        params.append(embed_model)
    rows = connection.execute(
        f"""
        SELECT *
        FROM messages
        WHERE duplicate_of_id IS NULL
          AND (
              enrichment_version < ?
              OR content_status IN ('pending', 'processing')
              {missing_vision}
              {missing_embedding}
          )
        ORDER BY
            CASE
                WHEN media_type = 'video'
                 AND datetime(date_utc) >= datetime('now', '-90 days')
                    THEN 0
                WHEN media_type = 'image'
                 AND datetime(date_utc) >= datetime('now', '-30 days')
                    THEN 1
                WHEN media_type = 'image'
                 AND (
                    LOWER(vision_text) LIKE '%person%'
                    OR LOWER(vision_text) LIKE '%people%'
                    OR LOWER(vision_text) LIKE '% man %'
                    OR LOWER(vision_text) LIKE '%woman%'
                    OR LOWER(vision_text) LIKE '%girl%'
                    OR LOWER(vision_text) LIKE '%boy%'
                    OR LOWER(vision_text) LIKE '%president%'
                    OR LOWER(vision_text) LIKE '%leader%'
                 ) THEN 2
                WHEN media_type = 'video' THEN 3
                WHEN media_type = 'image' THEN 4
                WHEN media_type IN ('pdf', 'document', 'audio') THEN 5
                WHEN urls_json != '[]' THEN 6
                ELSE 7
            END,
            date_utc DESC,
            message_id DESC
        LIMIT ?
        """,
        [*params, limit],
    ).fetchall()
    if rows:
        connection.executemany(
            "UPDATE messages SET content_status = 'processing' WHERE id = ?",
            [(row["id"],) for row in rows],
        )
        connection.commit()
    return [dict(row) for row in rows]
